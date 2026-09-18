"""Integração do fluxo real de importação, estudo e histórico."""
import io
import random
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import create_app
from simulados import choose_questions


def fixture_xlsx(bad=False, four_choices=False):
    wb = load_workbook(Path(__file__).resolve().parents[1] / "modelo_questoes.xlsx")
    ws = wb["Questoes"]
    values = dict(banca="Banca Exemplo", concurso="Concurso Fictício", cargo="Técnico",
                  ano=2023, data_prova="15/11/2023", prova="Prova A", numero="1", pagina=2,
                  materia_macro="Instrumentação", assunto_micro="Pressão", tipo="AE",
                  enunciado="Qual alternativa corresponde a 2 + 2?",
                  alternativa_a="1", alternativa_b="2", alternativa_c="3", alternativa_d="4",
                  alternativa_e="" if four_choices else "5", gabarito="D" if not bad else "Z", situacao="VALIDA",
                  explicacao="2 + 2 = 4.", fonte="Prova fictícia para teste")
    for idx, cell in enumerate(ws[1], 1):
        if cell.value in values:
            ws.cell(2, idx).value = values[cell.value]
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return data


class FlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({"TESTING": True, "DATABASE": str(Path(self.temp.name) / "test.sqlite3")})
        self.client = self.app.test_client()
        self.client.get("/")
        with self.client.session_transaction() as session:
            self.csrf = session["csrf"]

    def tearDown(self):
        self.temp.cleanup()

    def upload(self, stream):
        return self.client.post("/importar", data={"csrf": self.csrf, "file": (stream, "prova.xlsx")},
                                content_type="multipart/form-data", follow_redirects=True)

    def test_import_study_review_and_dashboard(self):
        self.assertIn(b"1 quest", self.upload(fixture_xlsx()).data)
        self.assertIn(b"1 quest", self.upload(fixture_xlsx()).data)  # sem duplicar
        r = self.client.get("/questoes?macro=Instrumenta%C3%A7%C3%A3o&mode=nao_vistas")
        self.assertIn(b"1 quest", r.data)
        r = self.client.post("/sessoes", data={"csrf": self.csrf, "macro": "Instrumentação",
                                                   "mode": "nao_vistas", "amount": "10"}, follow_redirects=True)
        self.assertIn("Qual alternativa".encode(), r.data)
        r = self.client.post("/sessoes/1/responder", data={"csrf": self.csrf, "answer": "A",
                                                           "seconds": "52", "doubt": "on"}, follow_redirects=True)
        self.assertIn("Vamos revisar".encode(), r.data)
        self.assertIn(b"Gabarito: D", r.data)
        self.assertIn("Pressão".encode(), self.client.get("/revisao").data)
        self.assertIn(b"52s", self.client.get("/dashboard").data)
        self.assertIn(b"0%", self.client.get("/sessoes/1/resultado").data)
        self.assertIn(b"0 quest", self.client.get("/questoes?mode=nao_vistas").data)
        self.assertIn(b"1 quest", self.client.get("/questoes?data_prova=2023-11-15").data)
        self.assertIn(b"0 quest", self.client.get("/questoes?data_prova=2023-11-16").data)
        self.assertIn(b"Imprimir / salvar PDF", self.client.get("/sessoes/1/imprimir").data)
        self.assertNotIn(b"Gabarito</h2>", self.client.get("/sessoes/1/imprimir").data)
        self.assertIn(b"Gabarito</h2>", self.client.get("/sessoes/1/imprimir?gabarito=1").data)
        self.assertEqual(self.client.post('/sessoes/1/editar', data={'csrf': self.csrf,
             'name': 'Metrologia teste'}).status_code, 302)
        self.assertIn(b'Metrologia teste', self.client.get('/historico?q=Metrologia').data)
        self.assertIn(b'Metrologia teste', self.client.get('/historico?inicio=2020-01-01&fim=2100-01-01').data)
        self.assertNotIn(b'Metrologia teste', self.client.get('/historico?fim=2020-01-01').data)
        self.assertEqual(self.client.post('/revisao/cartoes',data={'csrf':self.csrf,'front':'O que é pressão?',
                            'back':'Força por área.'},follow_redirects=True).status_code,200)
        self.assertIn('Força por área.'.encode(),self.client.get('/revisao').data)
        self.assertEqual(self.client.post('/revisao/cartoes/1/responder',data={'csrf':self.csrf,'grade':'good'}).status_code,302)
        self.assertIn(b'0 para revisar',self.client.get('/revisao').data)
        chart = self.client.get('/api/desempenho?period=7d').json
        self.assertEqual(chart['distribution'][0]['total'], 1)
        self.assertEqual(sum(chart['counts']), 1)

    def test_invalid_batch_leaves_database_empty_and_csrf_rejected(self):
        r = self.upload(fixture_xlsx(bad=True))
        self.assertIn("Importação cancelada".encode(), r.data)
        self.assertIn(b"0 quest", self.client.get("/questoes").data)
        self.assertEqual(self.client.post("/sessoes", data={"amount": "1"}).status_code, 400)

    def test_four_choices_and_import_guidance(self):
        model = self.client.get('/modelo-questoes')
        self.assertEqual(model.status_code, 200)
        self.assertIn(b'attachment', model.headers['Content-Disposition'].encode())
        model.close()
        self.assertIn(b'Copiar prompt', self.client.get('/importar').data)
        self.assertIn('56 respostas disponíveis'.encode(), self.client.get('/ajuda').data)
        self.assertIn(b'1 quest', self.upload(fixture_xlsx(four_choices=True)).data)
        self.client.post('/sessoes', data={'csrf':self.csrf,'macro':'Instrumentação','amount':'1'})
        page=self.client.get('/sessoes/1').data
        self.assertIn(b'value="D"',page)
        self.assertNotIn(b'value="E"',page)
        self.assertNotIn(b'(E)',self.client.get('/sessoes/1/imprimir').data)
        rejected=self.client.post('/sessoes/1/responder',data={'csrf':self.csrf,'answer':'E'},follow_redirects=True)
        self.assertIn('Selecione uma alternativa'.encode(),rejected.data)
        self.assertIn('Você acertou'.encode(),self.client.post('/sessoes/1/responder',data={'csrf':self.csrf,'answer':'D','seconds':'5'},follow_redirects=True).data)

    def test_real_proof_figure_filters_and_existing_classification_migration(self):
        source = Path(__file__).resolve().parents[1] / 'imports' / 'lote1_transpetro_petrobras.xlsx'
        if not source.is_file():
            self.skipTest('Acervo pessoal não incluído na versão pública')
        self.assertIn(b'711 novas', self.upload(source.open('rb')).data)
        with self.app.app_context():
            import sqlite3
            con = sqlite3.connect(self.app.config['DATABASE'])
            con.row_factory = sqlite3.Row
            target = con.execute("SELECT * FROM questions WHERE concurso='Transpetro 2011' AND numero='24'").fetchone()
            self.assertEqual(target['pagina'], '6')
            self.assertEqual(target['arquivo_pdf'], 'provas/T2011.pdf')
            sid = con.execute('INSERT INTO study_sessions DEFAULT VALUES').lastrowid
            con.execute('INSERT INTO session_questions VALUES (?,?,1)',(sid,target['id']))
            con.commit()
            con.close()
        study = self.client.get(f'/sessoes/{sid}').data
        self.assertIn(b'T2011_p006.webp', study)
        self.assertIn(b'Figura 1', study)
        self.assertIn(b'T2011_p006.webp', self.client.get(f'/sessoes/{sid}/imprimir').data)
        related=self.client.post(f'/revisao/semelhantes/{target["id"]}',data={'csrf':self.csrf})
        self.assertEqual(related.status_code,302)
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE'])
            self.assertGreater(con.execute("SELECT count(*) FROM session_questions WHERE session_id=2").fetchone()[0],0)
            con.close()
        self.assertIn(b'2011',self.client.get('/dashboard').data)
        self.assertGreater(len(self.client.get('/api/desempenho?macro=Desenho%20T%C3%A9cnico').json['micro_distribution']),0)
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE'])
            plain=con.execute("SELECT id FROM questions WHERE enunciado LIKE 'A projeção ortográfica consiste%' LIMIT 1").fetchone()
            con.close()
        if plain:
            with self.app.app_context():
                con=sqlite3.connect(self.app.config['DATABASE']);new_sid=con.execute('INSERT INTO study_sessions DEFAULT VALUES').lastrowid
                con.execute('INSERT INTO session_questions VALUES (?,?,1)',(new_sid,plain[0]));con.commit();con.close()
            page=self.client.get(f'/sessoes/{new_sid}').data
            self.assertIn(b'<details class="source-figure" >',page)
        result = self.client.get('/api/desempenho?period=total&concurso=Transpetro&ano=2011&banca=Cesgranrio').json
        self.assertEqual(sum(entry['total'] for entry in result['distribution']), 59)
        self.assertEqual(sum(entry['total'] for entry in self.client.get('/api/desempenho?concurso=Transpetro').json['distribution']),316)
        self.assertEqual(sum(result['counts']), 0)
        self.assertEqual(len(self.client.get('/api/desempenho?period=15d').json['labels']),15)
        self.client.post(f'/sessoes/{sid}/responder',data={'csrf':self.csrf,'answer':'A','seconds':'45'})
        self.assertEqual(sum(self.client.get('/api/desempenho?period=1d&concurso=Transpetro&ano=2011').json['counts']),1)
        self.assertEqual(sum(self.client.get('/api/desempenho?period=7d&concurso=Petrobras&ano=2011').json['counts']),0)
        self.assertIn(b'1 de 1 respondidas',self.client.get('/historico?status=concluida').data)
        con = sqlite3.connect(self.app.config['DATABASE'])
        con.execute("UPDATE questions SET macro='Conhecimentos Específicos',micro='A classificar' WHERE id=?",(target['id'],))
        con.commit();con.close()
        create_app({'TESTING':True,'DATABASE':self.app.config['DATABASE']})
        con = sqlite3.connect(self.app.config['DATABASE'])
        self.assertNotEqual(con.execute('SELECT micro FROM questions WHERE id=?',(target['id'],)).fetchone()[0], 'A classificar')
        self.assertEqual(con.execute('SELECT count(*) FROM session_questions WHERE session_id=?',(sid,)).fetchone()[0],1)
        con.close()
        self.assertTrue(Path(self.app.config['DATABASE']+'.backup-antes-da-classificacao').is_file())

    def test_simulator_hides_answers_until_final_and_keeps_exam_proportions(self):
        source=Path(__file__).resolve().parents[1] / 'imports' / 'lote1_transpetro_petrobras.xlsx'
        if not source.is_file():
            self.skipTest('Acervo pessoal não incluído na versão pública')
        self.upload(source.open('rb'))
        homepage=self.client.get('/simulados?concurso=Petrobras&quantidade=40')
        self.assertIn(b'40 quest',homepage.data)
        self.assertIn('Distribuição de referência'.encode(),homepage.data)
        start=self.client.post('/simulados',data={'csrf':self.csrf,'concurso':'Petrobras','quantidade':'40'})
        self.assertEqual(start.status_code,302)
        sid=int(start.headers['Location'].rsplit('/',1)[-1])
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE']);con.row_factory=sqlite3.Row
            chosen=con.execute('SELECT q.id,q.macro,q.micro,q.gabarito,q.banca,q.tipo FROM session_questions sq JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=? ORDER BY sq.position',(sid,)).fetchall()
            self.assertEqual(len(chosen),40)
            self.assertEqual(len({q['id'] for q in chosen}),40)
            self.assertGreaterEqual(sum(q['macro']=='Língua Portuguesa' for q in chosen),5)
            self.assertLessEqual(sum(q['macro']=='Informática' for q in chosen),3)
            con.close()
        self.assertIn(b'Tempo total',self.client.get(f'/sessoes/{sid}').data)
        self.assertEqual(self.client.get(f'/sessoes/{sid}/correcao/{chosen[0]["id"]}').status_code,302)
        self.assertEqual(self.client.get(f'/sessoes/{sid}/resultado').status_code,302)
        self.assertNotIn(b'Gabarito</h2>',self.client.get(f'/sessoes/{sid}/imprimir?gabarito=1').data)
        for i,q in enumerate(chosen):
            result=self.client.post(f'/sessoes/{sid}/responder',data={'csrf':self.csrf,'answer':'C' if q['tipo']=='CE' else 'A','seconds':'8',
                          'doubt':'on' if i==0 else ''})
            self.assertEqual(result.status_code,302)
            self.assertIn('/resultado' if i==39 else f'/sessoes/{sid}',result.headers['Location'])
            if i==0:
                self.assertIn(b'0</strong>',self.client.get('/dashboard').data)
                self.assertEqual(self.client.get(f'/sessoes/{sid}/correcao/{q["id"]}').status_code,302)
        result=self.client.get(f'/sessoes/{sid}/resultado').data
        self.assertIn('Desempenho por assunto'.encode(),result)
        self.assertIn(b'DURA',result)
        self.assertEqual(self.client.get(f'/sessoes/{sid}/correcao/{chosen[0]["id"]}').status_code,200)
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE'])
            self.assertEqual(con.execute('SELECT count(*) FROM attempts WHERE session_id=?',(sid,)).fetchone()[0],40)
            self.assertEqual(con.execute('SELECT count(*) FROM simulated_answers WHERE session_id=?',(sid,)).fetchone()[0],0)
            con.close()

    def test_micro_refinement_goal_and_sampler_with_exhausted_topic(self):
        source=Path(__file__).resolve().parents[1] / 'imports' / 'lote1_transpetro_petrobras.xlsx'
        if not source.is_file():
            self.skipTest('Acervo pessoal não incluído na versão pública')
        self.upload(source.open('rb'))
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE'])
            self.assertGreaterEqual(con.execute("SELECT count(DISTINCT micro) FROM questions WHERE macro='Matemática'").fetchone()[0],6)
            self.assertGreaterEqual(con.execute("SELECT count(DISTINCT micro) FROM questions WHERE macro='Língua Portuguesa'").fetchone()[0],6)
            target=con.execute("SELECT id FROM questions WHERE macro='Matemática' LIMIT 1").fetchone()[0]
            con.execute("UPDATE questions SET micro='Matemática' WHERE id=?",(target,))
            con.execute("UPDATE questions SET micro='Meu assunto manual' WHERE macro='Matemática' AND id<>? AND id=(SELECT min(id) FROM questions WHERE macro='Matemática' AND id<>?)",(target,target))
            con.commit();con.close()
        create_app({'TESTING':True,'DATABASE':self.app.config['DATABASE']})
        with self.app.app_context():
            import sqlite3
            con=sqlite3.connect(self.app.config['DATABASE'])
            self.assertNotEqual(con.execute('SELECT micro FROM questions WHERE id=?',(target,)).fetchone()[0],'Matemática')
            self.assertEqual(con.execute("SELECT count(*) FROM questions WHERE micro='Meu assunto manual'").fetchone()[0],1)
            con.close()
        self.assertEqual(self.client.post('/meta-concurso',data={'csrf':self.csrf,'exam_date':'2099-11-12','exam_title':'Transpetro'}).status_code,302)
        self.assertIn(b'2099-11-12',self.client.get('/dashboard').data)
        self.assertEqual(self.client.post('/meta-concurso',data={'csrf':self.csrf,'exam_date':''}).status_code,302)
        self.assertNotIn(b'2099-11-12',self.client.get('/dashboard').data)
        rows=[{'id':i,'macro':'Materiais','micro':'Metalografia','attempts':1,'mistakes':0,'doubts':0,'marked':0} for i in range(1,61)]
        rows += [{'id':i,'macro':'Informática','micro':'Planilhas','attempts':0,'mistakes':0,'doubts':0,'marked':0} for i in range(61,65)]
        sample=choose_questions(rows,30,random.Random(5))
        self.assertGreaterEqual(sum(i<=60 for i in sample),26)
        self.assertEqual(len(sample),len(set(sample)))


if __name__ == "__main__":
    unittest.main()
