"""Banco de questões local. Rode com: python app.py"""
import os
from collections import Counter
import secrets
import sqlite3
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import re

from flask import Flask, abort, flash, g, redirect, render_template, request, session, url_for, send_file
from openpyxl import load_workbook
from gerar_paginas import page_image
from assuntos import GENERIC, classify
from simulados import choose_questions, quota_preview
from help_content import HELP_GROUPS

ROOT = Path(__file__).resolve().parent
COLUMNS = [
    "banca", "concurso", "cargo", "ano", "data_prova", "prova", "numero", "pagina",
    "materia_macro", "assunto_micro", "tipo", "enunciado", "alternativa_a",
    "alternativa_b", "alternativa_c", "alternativa_d", "alternativa_e",
    "gabarito", "situacao", "explicacao", "fonte", "imagem", "arquivo_pdf",
]
REQUIRED = ["banca", "concurso", "cargo", "ano", "prova", "numero", "materia_macro",
            "assunto_micro", "enunciado", "gabarito", "fonte"]
VISUAL_REFERENCE = re.compile(r"\b(figura|figuras|tabela|tabelas|diagrama|esquema|imagem|croqui|ilustração|"
                              r"gráfico (?:abaixo|acima|a seguir|apresentado)|desenho (?:abaixo|acima|a seguir)|"
                              r"conforme (?:o desenho|a figura)|representado (?:na figura|abaixo))\b|fig\.", re.I)
CONTEST_NAME = "CASE WHEN substr(trim(q.concurso),-4)=CAST(q.ano AS TEXT) THEN trim(substr(trim(q.concurso),1,length(trim(q.concurso))-4)) ELSE trim(q.concurso) END"


def session_rows(conn, clause='', values=(), limit=8, offset=0):
    return conn.execute(f"""SELECT s.id,s.name,s.kind,s.started_at,s.finished_at,
         count(sq.question_id) total,count(coalesce(a.id,sa.question_id)) answered,coalesce(sum(a.correct),0) correct,
         coalesce(sum(a.seconds),0) study_seconds,
         coalesce(strftime('%s',coalesce(s.finished_at,max(a.answered_at),max(sa.answered_at),s.started_at))
                    - strftime('%s',s.started_at),0) elapsed_seconds,
         (SELECT group_concat(macro,' · ') FROM
             (SELECT DISTINCT q2.macro FROM session_questions sq2
              JOIN questions q2 ON q2.id=sq2.question_id
              WHERE sq2.session_id=s.id ORDER BY sq2.position LIMIT 2)) topics
         FROM study_sessions s JOIN session_questions sq ON sq.session_id=s.id
         LEFT JOIN attempts a ON a.session_id=sq.session_id AND a.question_id=sq.question_id
         LEFT JOIN simulated_answers sa ON sa.session_id=sq.session_id AND sa.question_id=sq.question_id
         WHERE 1=1 {clause}
         GROUP BY s.id ORDER BY s.id DESC LIMIT ? OFFSET ?""", (*values, limit, offset)).fetchall()


def period_start(code):
    days = {'1d': 0, '7d': 6, '15d': 14, '1m': 29, '3m': 89,
            '6m': 179, '1y': 364, 'total': None}
    if code not in days:
        return None
    return date.today() - timedelta(days=days[code]) if days[code] is not None else None


def create_app(test_config=None):
    app = Flask(__name__)
    app.config.update(DATABASE=str(ROOT / "instance" / "questoes.sqlite3"),
                      MAX_CONTENT_LENGTH=12 * 1024 * 1024)
    if test_config:
        app.config.update(test_config)
    dbpath = Path(app.config["DATABASE"])
    dbpath.parent.mkdir(parents=True, exist_ok=True)
    keypath = dbpath.parent / "secret.key"
    if app.config.get("TESTING"):
        app.secret_key = "only-for-tests"
    else:
        if not keypath.exists():
            keypath.write_text(secrets.token_hex(32), encoding="utf-8")
            try:
                os.chmod(keypath, 0o600)
            except OSError:
                pass
        app.secret_key = keypath.read_text(encoding="utf-8")

    def db():
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DATABASE"])
            g.db.row_factory = sqlite3.Row
            g.db.execute("PRAGMA foreign_keys = ON")
        return g.db

    @app.teardown_appcontext
    def close_db(error=None):
        connection = g.pop("db", None)
        if connection is not None:
            connection.close()

    with app.app_context():
        db().executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
        existing = {r[1] for r in db().execute("PRAGMA table_info(questions)")}
        if "tipo" not in existing:
            db().execute("ALTER TABLE questions ADD COLUMN tipo TEXT NOT NULL DEFAULT 'AE'")
        if "arquivo_pdf" not in existing:
            db().execute("ALTER TABLE questions ADD COLUMN arquivo_pdf TEXT NOT NULL DEFAULT ''")
        sessions_columns = {r[1] for r in db().execute("PRAGMA table_info(study_sessions)")}
        if "name" not in sessions_columns:
            db().execute("ALTER TABLE study_sessions ADD COLUMN name TEXT NOT NULL DEFAULT ''")
        db().commit()
        # Atualização pontual do primeiro lote, preservando categorias já editadas.
        classified = ROOT / 'imports' / 'lote1_transpetro_petrobras.xlsx'
        pending = db().execute("SELECT count(*) FROM questions WHERE macro='Conhecimentos Específicos' AND micro='A classificar'").fetchone()[0]
        if pending and classified.is_file():
            workbook = load_workbook(classified, read_only=True, data_only=True)
            try:
                rows = workbook['Questoes'].iter_rows(values_only=True)
                headers = [str(v).strip().lower() for v in next(rows)]
                idx = {field: headers.index(field) for field in ('banca','concurso','cargo','ano','prova','numero','materia_macro','assunto_micro')}
                classifications = []
                for cells in rows:
                    obj = {name: str(cells[pos]).strip() if cells[pos] is not None else '' for name,pos in idx.items()}
                    if not obj['banca'] or obj['assunto_micro'] == 'A classificar': continue
                    key = '|'.join(str(int(float(obj[name]))) if name == 'ano' else obj[name].casefold()
                                   for name in ('banca','concurso','cargo','ano','prova','numero'))
                    classifications.append((obj['materia_macro'],obj['assunto_micro'],key))
                if len(classifications) != 711:
                    raise ValueError('O lote de classificação deve conter 711 questões classificadas.')
            finally:
                workbook.close()
            backup = dbpath.with_name(dbpath.name + '.backup-antes-da-classificacao')
            if dbpath.is_file() and not backup.exists():
                with sqlite3.connect(backup) as dest:
                    db().backup(dest)
            with db():
                db().executemany("""UPDATE questions SET macro=?,micro=? WHERE source_key=?
                      AND macro='Conhecimentos Específicos' AND micro='A classificar'""", classifications)
        # Reclassifica somente os rótulos genéricos do acervo já importado.
        generic_rows = db().execute("SELECT id,macro,micro,enunciado FROM questions").fetchall()
        upgrades = [(classify(r['macro'],r['micro'],r['enunciado']),r['id'],r['micro'])
                    for r in generic_rows if r['micro'] in GENERIC.get(r['macro'],())]
        upgrades = [item for item in upgrades if item[0] != item[2]]
        if upgrades:
            backup = dbpath.with_name(dbpath.name + '.backup-antes-dos-assuntos-micro')
            if dbpath.is_file() and not backup.exists():
                with sqlite3.connect(backup) as dest:
                    db().backup(dest)
            with db():
                db().executemany('UPDATE questions SET micro=? WHERE id=? AND micro=?', upgrades)

    @app.before_request
    def csrf_check():
        if "csrf" not in session:
            session["csrf"] = secrets.token_hex(24)
        if request.method == "POST" and request.form.get("csrf") != session["csrf"]:
            abort(400, "Formulário expirado. Atualize a página e tente novamente.")

    @app.context_processor
    def common():
        def question_visual(q):
            if q['imagem'] and (ROOT / 'static' / q['imagem']).is_file():
                return url_for('static', filename=q['imagem'])
            path = page_image(q['arquivo_pdf'], q['pagina'])
            return url_for('static', filename=path) if path else None
        return {"csrf": session.get("csrf", ""), "question_visual": question_visual,
                "needs_figure": lambda q: bool(q['imagem']) or bool(VISUAL_REFERENCE.search(q['enunciado']))}

    @app.route("/")
    def home():
        conn = db()
        stats = conn.execute("""SELECT (SELECT count(*) FROM questions WHERE status='VALIDA') total,
                          (SELECT count(*) FROM attempts) attempts,
                          (SELECT count(*) FROM (SELECT DISTINCT question_id FROM attempts)) distinct_done,
                          (SELECT count(*) FROM attempts WHERE correct=1) correct""").fetchone()
        recent = session_rows(conn,limit=5)
        return render_template("home.html", stats=stats, recent=recent)

    @app.route('/historico')
    def history():
        conn = db()
        query = request.args.get('q', '').strip()[:80]
        status = request.args.get('status', '')
        macro = request.args.get('macro', '')
        start_date = request.args.get('inicio','')
        end_date = request.args.get('fim','')
        for value in (start_date,end_date):
            if value:
                try: date.fromisoformat(value)
                except ValueError: abort(400)
        page = max(1, min(100000, request.args.get('page', 1, type=int)))
        clause, values = '', []
        if query:
            clause += ' AND (s.name LIKE ? OR CAST(s.id AS TEXT)=? OR EXISTS (SELECT 1 FROM session_questions sqx JOIN questions qx ON qx.id=sqx.question_id WHERE sqx.session_id=s.id AND (qx.macro LIKE ? OR qx.micro LIKE ?)))'
            values += [f'%{query}%', query, f'%{query}%', f'%{query}%']
        if status == 'concluida': clause += ' AND s.finished_at IS NOT NULL'
        if status == 'andamento': clause += ' AND s.finished_at IS NULL'
        if macro:
            clause += ' AND EXISTS (SELECT 1 FROM session_questions sqx JOIN questions qx ON qx.id=sqx.question_id WHERE sqx.session_id=s.id AND qx.macro=?)'
            values.append(macro)
        if start_date:
            clause += ' AND date(s.started_at,\'localtime\')>=?'
            values.append(start_date)
        if end_date:
            clause += ' AND date(s.started_at,\'localtime\')<=?'
            values.append(end_date)
        total = conn.execute(f'SELECT count(*) FROM study_sessions s WHERE 1=1 {clause}', values).fetchone()[0]
        rows = session_rows(conn, clause, values, limit=20, offset=(page-1)*20)
        macros = [r[0] for r in conn.execute('SELECT DISTINCT macro FROM questions ORDER BY macro')]
        return render_template('history.html', rows=rows,total=total,page=page,
                               query=query,status=status,macro=macro,macros=macros,
                               start_date=start_date,end_date=end_date)

    @app.post('/sessoes/<int:sid>/editar')
    def edit_session(sid):
        title = request.form.get('name','').strip()
        if len(title)>80:
            abort(400)
        with db():
            cursor = db().execute('UPDATE study_sessions SET name=? WHERE id=?',(title,sid))
        if not cursor.rowcount: abort(404)
        flash('Nome da sessão salvo.','success')
        return redirect(url_for('history'))

    @app.route("/questoes")
    def catalog():
        conn = db()
        filters = {k: request.args.get(k, "").strip() for k in ("macro", "micro", "banca", "ano", "data_prova", "cargo")}
        wheres = ["q.status='VALIDA'"]
        vals = []
        for field in ("macro", "micro", "banca", "ano", "data_prova", "cargo"):
            if filters[field]:
                wheres.append(f"q.{field}=?")
                vals.append(filters[field])
        mode = request.args.get("mode", "todas")
        if mode == "nao_vistas":
            wheres.append("NOT EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id)")
        elif mode == "erradas":
            wheres.append("EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id AND a.correct=0)")
        elif mode == "marcadas":
            wheres.append("EXISTS (SELECT 1 FROM review_flags f WHERE f.question_id=q.id AND f.marked=1)")
        elif mode == "duvida":
            wheres.append("EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id AND a.doubt=1)")
        elif mode != "todas":
            abort(400)
        where = " AND ".join(wheres)
        count = conn.execute(f"SELECT count(*) FROM questions q WHERE {where}", vals).fetchone()[0]
        preview = conn.execute(f"SELECT q.* FROM questions q WHERE {where} ORDER BY q.id DESC LIMIT 12", vals).fetchall()
        options = {field: [r[0] for r in conn.execute(
            f"SELECT DISTINCT {field} FROM questions WHERE status='VALIDA' ORDER BY {field}")]
            for field in ("macro", "micro", "banca", "ano", "data_prova", "cargo")}
        topic_map = {}
        for r in conn.execute("SELECT DISTINCT macro,micro FROM questions WHERE status='VALIDA' ORDER BY macro,micro"):
            topic_map.setdefault(r['macro'], []).append(r['micro'])
        return render_template("catalog.html", filters=filters, options=options,
                               topic_map=topic_map,mode=mode, count=count, preview=preview)

    @app.post("/sessoes")
    def new_session():
        conn = db()
        allowed = ("macro", "micro", "banca", "ano", "data_prova", "cargo")
        wheres = ["q.status='VALIDA'"]
        vals = []
        for field in allowed:
            value = request.form.get(field, "").strip()
            if value:
                wheres.append(f"q.{field}=?")
                vals.append(value)
        mode = request.form.get("mode", "todas")
        clauses = {"nao_vistas": "NOT EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id)",
                   "erradas": "EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id AND a.correct=0)",
                   "marcadas": "EXISTS (SELECT 1 FROM review_flags f WHERE f.question_id=q.id AND f.marked=1)",
                   "duvida": "EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id AND a.doubt=1)"}
        if mode != "todas" and mode not in clauses:
            abort(400)
        if mode in clauses:
            wheres.append(clauses[mode])
        try:
            amount = int(request.form.get("amount", "10"))
        except ValueError:
            abort(400)
        if not 1 <= amount <= 100:
            abort(400)
        order = "RANDOM()" if request.form.get("shuffle") == "on" else "q.id"
        ids = [r[0] for r in conn.execute(
            f"SELECT q.id FROM questions q WHERE {' AND '.join(wheres)} ORDER BY {order} LIMIT ?", vals + [amount])]
        if not ids:
            flash("Nenhuma questão encontrada para esses filtros.", "warning")
            return redirect(url_for("catalog"))
        with conn:
            cursor = conn.execute("INSERT INTO study_sessions DEFAULT VALUES")
            sid = cursor.lastrowid
            conn.executemany("INSERT INTO session_questions VALUES (?,?,?)",
                             [(sid, qid, pos) for pos, qid in enumerate(ids, 1)])
        return redirect(url_for("study", sid=sid))

    def simulator_pool(contest, board=''):
        where=["q.status='VALIDA'",f'{CONTEST_NAME}=?']
        args=[contest]
        if board:
            where.append('q.banca=?');args.append(board)
        return db().execute(f"""SELECT q.id,q.macro,q.micro,count(a.id) attempts,
                coalesce(sum(CASE WHEN a.correct=0 THEN 1 ELSE 0 END),0) mistakes,
                coalesce(sum(a.doubt),0) doubts,coalesce(f.marked,0) marked
                FROM questions q LEFT JOIN attempts a ON a.question_id=q.id
                LEFT JOIN review_flags f ON f.question_id=q.id
                WHERE {' AND '.join(where)} GROUP BY q.id ORDER BY q.id""",args).fetchall()

    @app.get('/simulados')
    def simulator():
        conn=db()
        contests=[r[0] for r in conn.execute(f"SELECT DISTINCT {CONTEST_NAME} FROM questions q WHERE q.status='VALIDA' ORDER BY 1")]
        boards=[r[0] for r in conn.execute("SELECT DISTINCT banca FROM questions WHERE status='VALIDA' ORDER BY banca")]
        contest=request.args.get('concurso','').strip()
        board=request.args.get('banca','').strip()
        amount=request.args.get('quantidade',60,type=int)
        amount=max(5,min(150,amount))
        if contest not in contests: contest=''
        if board not in boards: board=''
        pool=simulator_pool(contest,board) if contest else []
        counts=quota_preview(pool,min(amount,len(pool)))
        return render_template('simulators.html',contests=contests,boards=boards,
                               contest=contest,board=board,amount=amount,
                               available=len(pool),distribution=counts)

    @app.post('/simulados')
    def new_simulator():
        contest=request.form.get('concurso','').strip()
        board=request.form.get('banca','').strip()
        try: amount=int(request.form.get('quantidade','60'))
        except ValueError: abort(400)
        if not 5<=amount<=150 or not contest: abort(400)
        pool=simulator_pool(contest,board)
        if len(pool)<5:
            flash('Não há questões suficientes para esse concurso e banca.','warning')
            return redirect(url_for('simulator',concurso=contest,banca=board,quantidade=amount))
        if len(pool)<amount:
            flash(f'O acervo tem {len(pool)} questões para esses filtros; o simulado usará todas.','warning')
        chosen=choose_questions(pool,amount)
        with db():
            cursor=db().execute("INSERT INTO study_sessions(name,kind) VALUES (?, 'simulado')",
                                (f'Simulado {contest}' + (f' · {board}' if board else ''),))
            sid=cursor.lastrowid
            db().executemany('INSERT INTO session_questions VALUES (?,?,?)',
                             [(sid,qid,i) for i,qid in enumerate(chosen,1)])
        return redirect(url_for('study',sid=sid))

    def session_state(sid):
        conn = db()
        state = conn.execute("SELECT * FROM study_sessions WHERE id=?", (sid,)).fetchone()
        if not state:
            abort(404)
        rows = conn.execute("""SELECT q.*,sq.position,coalesce(a.answer,sa.answer) answer,a.correct,
                          coalesce(a.seconds,sa.seconds) seconds,coalesce(a.doubt,sa.doubt) doubt,
                          coalesce(f.marked,0) marked FROM session_questions sq
                          JOIN questions q ON q.id=sq.question_id
                          LEFT JOIN attempts a ON a.session_id=sq.session_id AND a.question_id=q.id
                          LEFT JOIN simulated_answers sa ON sa.session_id=sq.session_id AND sa.question_id=q.id
                          LEFT JOIN review_flags f ON f.question_id=q.id
                          WHERE sq.session_id=? ORDER BY sq.position""", (sid,)).fetchall()
        return state, rows

    @app.route("/sessoes/<int:sid>")
    def study(sid):
        state, rows = session_state(sid)
        pending = next((r for r in rows if r["answer"] is None), None)
        if pending is None:
            if not state["finished_at"]:
                db().execute("UPDATE study_sessions SET finished_at=CURRENT_TIMESTAMP WHERE id=?", (sid,))
                db().commit()
            return redirect(url_for("session_result", sid=sid))
        return render_template("study.html", sid=sid, state=state, question=pending, rows=rows,
                               answered=sum(r["answer"] is not None for r in rows),
                               macro_count=len({r['macro'] for r in rows}),
                               micro_count=len({(r['macro'],r['micro']) for r in rows}))

    @app.post("/sessoes/<int:sid>/responder")
    def answer(sid):
        state, rows = session_state(sid)
        pending = next((r for r in rows if r["answer"] is None), None)
        if pending is None:
            return redirect(url_for("session_result", sid=sid))
        choice = request.form.get("answer", "").upper()
        choices = "CE" if pending["tipo"] == "CE" else ("ABCDE" if pending["alternativa_e"] else "ABCD")
        if choice not in choices or len(choice) != 1:
            flash("Selecione uma alternativa.", "warning")
            return redirect(url_for("study", sid=sid))
        try:
            seconds = max(0, min(86400, int(request.form.get("seconds", "0"))))
        except ValueError:
            seconds = 0
        conn = db()
        with conn:
            if state['kind']=='simulado':
                conn.execute("""INSERT INTO simulated_answers(session_id,question_id,answer,seconds,doubt)
                                VALUES (?,?,?,?,?)""",(sid,pending['id'],choice,seconds,
                                int(request.form.get('doubt')=='on')))
                if sum(r['answer'] is None for r in rows)==1:
                    conn.execute("""INSERT INTO attempts(session_id,question_id,answer,correct,seconds,doubt,answered_at)
                         SELECT sa.session_id,sa.question_id,sa.answer,CAST(sa.answer=q.gabarito AS INTEGER),
                                sa.seconds,sa.doubt,sa.answered_at
                         FROM simulated_answers sa JOIN questions q ON q.id=sa.question_id
                         WHERE sa.session_id=?""",(sid,))
                    conn.execute('DELETE FROM simulated_answers WHERE session_id=?',(sid,))
                    conn.execute('UPDATE study_sessions SET finished_at=CURRENT_TIMESTAMP WHERE id=?',(sid,))
                return redirect(url_for('session_result' if sum(r['answer'] is None for r in rows)==1 else 'study',sid=sid))
            conn.execute("""INSERT INTO attempts(session_id,question_id,answer,correct,seconds,doubt)
                            VALUES (?,?,?,?,?,?)""", (sid, pending["id"], choice,
                            int(choice == pending["gabarito"]), seconds,
                            int(request.form.get("doubt") == "on")))
            if sum(r['answer'] is None for r in rows) == 1:
                conn.execute('UPDATE study_sessions SET finished_at=CURRENT_TIMESTAMP WHERE id=?', (sid,))
        return redirect(url_for("feedback", sid=sid, qid=pending["id"])+'#correcao')

    @app.route("/sessoes/<int:sid>/correcao/<int:qid>")
    def feedback(sid, qid):
        state, rows = session_state(sid)
        if state['kind']=='simulado' and not state['finished_at']:
            return redirect(url_for('study',sid=sid))
        question = next((r for r in rows if r["id"] == qid and r["answer"] is not None), None)
        if not question:
            abort(404)
        pending = any(r["answer"] is None for r in rows)
        return render_template("feedback.html", question=question, sid=sid, pending=pending)

    @app.route("/sessoes/<int:sid>/resultado")
    def session_result(sid):
        state, rows = session_state(sid)
        if state['kind']=='simulado' and not state['finished_at']:
            return redirect(url_for('study',sid=sid))
        done = [r for r in rows if r["answer"] is not None]
        groups=Counter((r['macro'],r['micro']) for r in rows)
        results=Counter((r['macro'],r['micro']) for r in done if r['correct'])
        seconds=max(0,db().execute("SELECT strftime('%s',finished_at)-strftime('%s',started_at) FROM study_sessions WHERE id=?",(sid,)).fetchone()[0] or 0)
        return render_template("result.html", sid=sid, rows=rows, done=done,
                               correct=sum(r["correct"] for r in done),state=state,
                               elapsed_seconds=seconds,topic_results=[(m,t,n,results[m,t])
                                   for (m,t),n in sorted(groups.items(),key=lambda item:(-item[1],item[0]))])

    @app.route('/sessoes/<int:sid>/imprimir')
    def print_session(sid):
        state, rows = session_state(sid)
        return render_template('print_session.html', state=state, rows=rows,
                               show_answers=request.args.get('gabarito') == '1' and
                               (state['kind']!='simulado' or bool(state['finished_at'])))

    @app.post("/questoes/<int:qid>/marcar")
    def toggle_flag(qid):
        conn = db()
        if not conn.execute("SELECT id FROM questions WHERE id=?", (qid,)).fetchone():
            abort(404)
        with conn:
            conn.execute("""INSERT INTO review_flags(question_id,marked) VALUES (?,1)
                            ON CONFLICT(question_id) DO UPDATE SET marked=1-marked""", (qid,))
        target = request.form.get("next", "")
        if not target.startswith("/") or target.startswith("//"):
            target = url_for("review")
        return redirect(target)

    @app.route("/revisao")
    def review():
        conn = db()
        rows = conn.execute("""SELECT q.*,coalesce(f.marked,0) marked,
             (SELECT count(*) FROM attempts a WHERE a.question_id=q.id AND a.correct=0) mistakes,
             (SELECT count(*) FROM attempts a WHERE a.question_id=q.id AND a.doubt=1) doubts,
             (SELECT a.correct FROM attempts a WHERE a.question_id=q.id ORDER BY a.id DESC LIMIT 1) last_correct
             FROM questions q LEFT JOIN review_flags f ON f.question_id=q.id
             WHERE q.status='VALIDA' AND (coalesce(f.marked,0)=1 OR
               EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id AND (a.correct=0 OR a.doubt=1)))
             ORDER BY CASE WHEN (SELECT a.correct FROM attempts a WHERE a.question_id=q.id ORDER BY a.id DESC LIMIT 1)=0
                       THEN 0 ELSE 1 END, q.id DESC""").fetchall()
        recommendations = {}
        for q in rows[:30]:
            related = conn.execute("""SELECT q.id,q.numero,q.ano,q.banca,q.micro,q.enunciado,
                     CASE WHEN q.micro=? THEN 0 ELSE 1 END relevance
                     FROM questions q WHERE q.status='VALIDA' AND q.id<>? AND q.macro=?
                     AND NOT EXISTS (SELECT 1 FROM attempts a WHERE a.question_id=q.id)
                     ORDER BY relevance,q.id LIMIT 3""",(q['micro'],q['id'],q['macro'])).fetchall()
            recommendations[q['id']] = related
        due = conn.execute("SELECT count(*) FROM flashcards WHERE due_at<=CURRENT_TIMESTAMP").fetchone()[0]
        cards = conn.execute("""SELECT f.*,q.micro FROM flashcards f LEFT JOIN questions q ON q.id=f.question_id
                              ORDER BY (f.due_at<=CURRENT_TIMESTAMP) DESC,f.due_at,f.id LIMIT 40""").fetchall()
        return render_template("review.html", rows=rows, recommendations=recommendations,cards=cards,due=due,
                               now=datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S'))

    @app.post('/revisao/semelhantes/<int:qid>')
    def similar_session(qid):
        conn=db()
        original=conn.execute("SELECT * FROM questions WHERE id=? AND status='VALIDA'",(qid,)).fetchone()
        if not original: abort(404)
        related=conn.execute("""SELECT q.id FROM questions q WHERE q.status='VALIDA' AND q.id<>?
            AND q.macro=? AND NOT EXISTS(SELECT 1 FROM attempts a WHERE a.question_id=q.id)
            ORDER BY CASE WHEN q.micro=? THEN 0 ELSE 1 END,RANDOM() LIMIT 10""",
            (qid,original['macro'],original['micro'])).fetchall()
        if not related:
            flash('Não há questões novas desta matéria disponíveis.','warning')
            return redirect(url_for('review'))
        with conn:
            cursor=conn.execute("INSERT INTO study_sessions(name,kind) VALUES (?,?)",
                 (f"Praticar {original['micro'][:60]}",'revisao'))
            conn.executemany("INSERT INTO session_questions VALUES (?,?,?)",
                [(cursor.lastrowid,row['id'],i) for i,row in enumerate(related,1)])
        return redirect(url_for('study',sid=cursor.lastrowid))

    @app.post('/revisao/cartoes')
    def add_card():
        front=request.form.get('front','').strip()[:2000]
        back=request.form.get('back','').strip()[:4000]
        qid=request.form.get('question_id',type=int)
        if not front or not back: abort(400)
        if qid and not db().execute('SELECT 1 FROM questions WHERE id=?',(qid,)).fetchone(): abort(400)
        with db():
            db().execute('INSERT INTO flashcards(question_id,front,back) VALUES (?,?,?)',(qid,front,back))
        flash('Cartão criado para revisão.','success')
        return redirect(url_for('review')+'#cartoes')

    @app.post('/revisao/cartoes/<int:cid>/responder')
    def grade_card(cid):
        grade=request.form.get('grade')
        if grade not in ('again','good','easy'): abort(400)
        conn=db();card=conn.execute('SELECT * FROM flashcards WHERE id=?',(cid,)).fetchone()
        if not card: abort(404)
        if grade=='again': days=0; repetitions=0
        elif grade=='easy': days=min(180,max(4,card['interval_days']*3));repetitions=card['repetitions']+1
        else: days=min(180,max(1,card['interval_days']*2));repetitions=card['repetitions']+1
        due=(datetime.now(timezone.utc)+timedelta(days=days,minutes=10 if grade=='again' else 0)).strftime('%Y-%m-%d %H:%M:%S')
        with conn:
            conn.execute('UPDATE flashcards SET due_at=?,interval_days=?,repetitions=?,reviewed_at=CURRENT_TIMESTAMP WHERE id=?',
                         (due,days,repetitions,cid))
        return redirect(url_for('review')+'#cartoes')

    @app.post('/revisao/cartoes/<int:cid>/excluir')
    def delete_card(cid):
        with db(): db().execute('DELETE FROM flashcards WHERE id=?',(cid,))
        return redirect(url_for('review')+'#cartoes')

    @app.route("/dashboard")
    def dashboard():
        conn = db()
        overall = conn.execute("""SELECT count(*) total,coalesce(sum(correct),0) correct,
                coalesce(round(avg(seconds)),0) avg_seconds,count(DISTINCT question_id) unique_done
                FROM attempts""").fetchone()
        topics = conn.execute("""SELECT q.macro,q.micro,count(a.id) total,sum(a.correct) correct,
               round(avg(a.seconds)) avg_seconds,count(DISTINCT q.id) unique_done
               FROM attempts a JOIN questions q ON q.id=a.question_id
               GROUP BY q.macro,q.micro ORDER BY total DESC,q.macro,q.micro""").fetchall()
        daily = conn.execute("""SELECT date(answered_at) day,count(*) total,sum(correct) correct
                              FROM attempts GROUP BY date(answered_at) ORDER BY day DESC LIMIT 14""").fetchall()
        contests = [r[0] for r in conn.execute(f"SELECT DISTINCT {CONTEST_NAME} FROM questions q WHERE q.status='VALIDA' ORDER BY 1")]
        years = [r[0] for r in conn.execute("SELECT DISTINCT ano FROM questions WHERE status='VALIDA' ORDER BY ano DESC")]
        boards = [r[0] for r in conn.execute("SELECT DISTINCT banca FROM questions WHERE status='VALIDA' ORDER BY banca")]
        unseen_count=conn.execute("SELECT count(*) FROM questions q WHERE status='VALIDA' AND NOT EXISTS(SELECT 1 FROM attempts a WHERE a.question_id=q.id)").fetchone()[0]
        review_count=conn.execute("""SELECT count(*) FROM questions q LEFT JOIN review_flags f ON f.question_id=q.id
            WHERE q.status='VALIDA' AND (coalesce(f.marked,0)=1 OR EXISTS(SELECT 1 FROM attempts a WHERE a.question_id=q.id AND (a.correct=0 OR a.doubt=1)))""").fetchone()[0]
        due_cards=conn.execute('SELECT count(*) FROM flashcards WHERE due_at<=CURRENT_TIMESTAMP').fetchone()[0]
        goal=dict(conn.execute("SELECT key,value FROM settings WHERE key IN ('exam_date','exam_title')").fetchall())
        try: days_remaining=(date.fromisoformat(goal['exam_date'])-date.today()).days if goal.get('exam_date') else None
        except ValueError: days_remaining=None
        return render_template("dashboard.html", overall=overall, topics=topics,
                               daily=list(reversed(daily)), contests=contests, boards=boards,years=years,
                               unseen_count=unseen_count,review_count=review_count,due_cards=due_cards,
                               exam_goal=goal,days_remaining=days_remaining)

    @app.post('/meta-concurso')
    def set_exam_goal():
        raw=request.form.get('exam_date','').strip()
        title=request.form.get('exam_title','').strip()[:65]
        if raw:
            try: day=date.fromisoformat(raw)
            except ValueError: abort(400)
            if not 2000<=day.year<=2100: abort(400)
        with db():
            if raw:
                db().executemany("""INSERT INTO settings(key,value) VALUES (?,?)
                    ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
                    [('exam_date',raw),('exam_title',title)])
            else:
                db().execute("DELETE FROM settings WHERE key IN ('exam_date','exam_title')")
        flash('Data do concurso atualizada.' if raw else 'Data do concurso removida.','success')
        return redirect(url_for('dashboard'))

    @app.get('/api/desempenho')
    def chart_data():
        conn = db()
        period = request.args.get('period','7d')
        if period not in ('1d','7d','15d','1m','3m','6m','1y','total'): abort(400)
        contest = request.args.get('concurso','').strip()
        board = request.args.get('banca','').strip()
        year = request.args.get('ano','').strip()
        if year and (not year.isdigit() or len(year)!=4): abort(400)
        selected_macro=request.args.get('macro','').strip()
        source_where, source_args = ["q.status='VALIDA'"], []
        if contest: source_where.append(f'{CONTEST_NAME}=?');source_args.append(contest)
        if board: source_where.append('q.banca=?');source_args.append(board)
        if year: source_where.append('q.ano=?');source_args.append(int(year))
        distribution = conn.execute(f"SELECT q.macro label,count(*) total FROM questions q WHERE {' AND '.join(source_where)} GROUP BY q.macro ORDER BY total DESC,q.macro",source_args).fetchall()
        micro_distribution=[]
        if selected_macro:
            micro_distribution=conn.execute(f"SELECT q.micro label,count(*) total FROM questions q WHERE {' AND '.join(source_where)} AND q.macro=? GROUP BY q.micro ORDER BY total DESC,q.micro",(*source_args,selected_macro)).fetchall()
        attempt_where = source_where.copy()
        start = period_start(period)
        if start: attempt_where.append("date(a.answered_at,'localtime')>=?");source_args.append(start.isoformat())
        if period == '1d':
            group = "strftime('%H',a.answered_at,'localtime')"
        elif period == 'total':
            first = conn.execute(f"SELECT min(date(a.answered_at,'localtime')) FROM attempts a JOIN questions q ON q.id=a.question_id WHERE {' AND '.join(attempt_where)}", source_args).fetchone()[0]
            group = "strftime('%Y-%m',a.answered_at,'localtime')" if first and (date.today()-date.fromisoformat(first)).days>366 else "date(a.answered_at,'localtime')"
        else:
            group = "date(a.answered_at,'localtime')"
        data = conn.execute(f"SELECT {group} bucket,count(*) total,sum(a.correct) correct FROM attempts a JOIN questions q ON q.id=a.question_id WHERE {' AND '.join(attempt_where)} GROUP BY bucket ORDER BY bucket",source_args).fetchall()
        lookup = {row['bucket']:(row['total'],row['correct']) for row in data}
        if period == '1d':
            labels = [f'{h:02d}h' for h in range(datetime.now().hour+1)]
            results = [lookup.get(f'{h:02d}',(0,0)) for h in range(len(labels))]
        elif period == 'total' and group.startswith("strftime"):
            if data:
                first = date.fromisoformat(data[0]['bucket']+'-01')
                labels=[];year,month=first.year,first.month
                while (year,month)<=(date.today().year,date.today().month):
                    labels.append(f'{year}-{month:02d}')
                    month+=1
                    if month==13:year+=1;month=1
            else: labels=[]
            results=[lookup.get(label,(0,0)) for label in labels]
        else:
            first = start or (date.fromisoformat(data[0]['bucket']) if data else date.today())
            labels=[(first+timedelta(days=d)).isoformat() for d in range((date.today()-first).days+1)]
            results=[lookup.get(label,(0,0)) for label in labels]
        return {'labels':labels,'counts':[r[0] for r in results],
                'correct':[r[1] for r in results],
                'distribution':[dict(r) for r in distribution],
                'micro_distribution':[dict(r) for r in micro_distribution],
                'period':period,'grouped_monthly':period=='total' and group.startswith('strftime') and period!='1d'}

    @app.route("/importar", methods=["GET", "POST"])
    def import_page():
        if request.method == "GET":
            return render_template("import.html", prompt_text=(ROOT / "docs" / "prompt-importacao.txt").read_text(encoding="utf-8"))
        file = request.files.get("file")
        if not file or not file.filename.lower().endswith(".xlsx"):
            flash("Selecione uma planilha .xlsx.", "warning")
            return redirect(url_for("import_page"))
        try:
            result = import_workbook(db(), file.stream)
        except (ValueError, OSError, KeyError) as exc:
            flash(str(exc), "warning")
            return redirect(url_for("import_page"))
        flash(f"Importação concluída: {result['created']} novas, {result['duplicates']} já existentes.", "success")
        return redirect(url_for("catalog"))

    @app.get("/modelo-questoes")
    def download_model():
        return send_file(ROOT / "modelo_questoes.xlsx", as_attachment=True,
                         download_name="modelo_questoes_caderno_local.xlsx")

    @app.get("/ajuda")
    def help_page():
        return render_template("help.html", groups=HELP_GROUPS,
                               question_count=sum(len(group["items"]) for group in HELP_GROUPS))

    return app


def import_workbook(conn, stream):
    """Valida o lote inteiro antes de gravar; uma linha inválida bloqueia o arquivo."""
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Não foi possível ler o arquivo .xlsx.") from exc
    try:
        if "Questoes" not in wb.sheetnames:
            raise ValueError("A planilha precisa ter uma aba chamada 'Questoes'.")
        sheet = wb["Questoes"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v).strip().lower() if v is not None else "" for v in next(rows, ())]
        # Versões anteriores do modelo continuam aceitas sem tipo/PDF.
        missing = set(COLUMNS) - {"tipo", "arquivo_pdf"} - set(headers)
        if missing:
            raise ValueError("Faltam colunas: " + ", ".join(sorted(missing)))
        indexes = {col: headers.index(col) if col in headers else -1 for col in COLUMNS}
        parsed, errors, seen = [], [], set()
        for rownum, cells in enumerate(rows, 2):
            if not any(v is not None and str(v).strip() for v in cells):
                continue
            obj = {col: str(cells[indexes[col]]).strip() if indexes[col]>=0 and indexes[col] < len(cells) and
                   cells[indexes[col]] is not None else "" for col in COLUMNS}
            obj['assunto_micro'] = classify(obj['materia_macro'],obj['assunto_micro'],obj['enunciado'])
            absent = [col for col in REQUIRED if not obj[col]]
            if absent:
                errors.append(f"linha {rownum}: campos vazios: {', '.join(absent)}")
                continue
            kind = obj["tipo"].upper() or "AE"
            if kind not in ("AE", "CE"):
                errors.append(f"linha {rownum}: tipo deve ser AE ou CE")
                continue
            choice_cols = "ce" if kind == "CE" else "abcd"
            if any(not obj[f"alternativa_{l}"] for l in choice_cols):
                errors.append(f"linha {rownum}: alternativas incompletas para {kind}")
                continue
            try:
                year = int(float(obj["ano"]))
                if str(year) != obj["ano"] and str(float(year)) != obj["ano"]:
                    raise ValueError
                if not 1950 <= year <= 2100:
                    raise ValueError
            except ValueError:
                errors.append(f"linha {rownum}: ano inválido")
                continue
            raw_date = obj["data_prova"]
            if raw_date:
                try:
                    obj["data_prova"] = datetime.fromisoformat(raw_date).date().isoformat()
                except ValueError:
                    try:
                        obj["data_prova"] = datetime.strptime(raw_date, "%d/%m/%Y").date().isoformat()
                    except ValueError:
                        errors.append(f"linha {rownum}: data_prova inválida (use DD/MM/AAAA ou AAAA-MM-DD)")
            key = "|".join((str(year) if k == "ano" else obj[k].casefold().strip()
                            for k in ("banca", "concurso", "cargo", "ano", "prova", "numero")))
            if key in seen:
                errors.append(f"linha {rownum}: identificação repetida dentro da planilha")
            seen.add(key)
            answer = obj["gabarito"].upper()
            status = obj["situacao"].upper() or "VALIDA"
            if answer not in ("A", "B", "C", "D", "E", "ANULADA") or status not in ("VALIDA", "ANULADA") or (answer == "ANULADA") != (status == "ANULADA"):
                errors.append(f"linha {rownum}: gabarito/situação inconsistentes")
            elif status == "VALIDA" and (answer not in ("C", "E") if kind == "CE" else not obj[f"alternativa_{answer.lower()}"]):
                errors.append(f"linha {rownum}: gabarito não corresponde a uma alternativa preenchida")
            image = obj["imagem"].replace("\\", "/")
            if image and (image.startswith("/") or ".." in Path(image).parts or not image.startswith("imagens/")):
                errors.append(f"linha {rownum}: imagem deve usar caminho relativo imagens/nome.ext")
            pdf = obj["arquivo_pdf"].replace("\\", "/")
            if pdf and (pdf.startswith("/") or ".." in Path(pdf).parts or not pdf.startswith("provas/") or not pdf.lower().endswith(".pdf")):
                errors.append(f"linha {rownum}: arquivo_pdf deve usar caminho relativo provas/nome.pdf")
            parsed.append((key, obj, year, answer, status, kind, image, pdf))
        if errors:
            raise ValueError("Importação cancelada; corrija: " + "; ".join(errors[:8]) +
                             (f"; e mais {len(errors)-8} erros" if len(errors)>8 else ""))
        if not parsed:
            raise ValueError("A aba Questoes não contém questões preenchidas.")
        created = duplicates = 0
        with conn:
            for key, obj, year, answer, status, kind, image, pdf in parsed:
                values = (key, obj["banca"], obj["concurso"], obj["cargo"], year, obj["data_prova"],
                          obj["prova"], obj["numero"], obj["pagina"], obj["materia_macro"],
                          obj["assunto_micro"], kind, obj["enunciado"], *(obj[f"alternativa_{l}"] for l in "abcde"),
                          answer, status, obj["explicacao"], obj["fonte"], image, pdf)
                cursor = conn.execute("""INSERT OR IGNORE INTO questions
                  (source_key,banca,concurso,cargo,ano,data_prova,prova,numero,pagina,macro,micro,tipo,enunciado,
                  alternativa_a,alternativa_b,alternativa_c,alternativa_d,alternativa_e,
                  gabarito,status,explicacao,fonte,imagem,arquivo_pdf)
                  VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", values)
                created += cursor.rowcount
                duplicates += int(cursor.rowcount == 0)
        return {"created": created, "duplicates": duplicates}
    finally:
        wb.close()


if __name__ == "__main__":
    create_app().run(host="127.0.0.1", port=5000, debug=False)
