"""Gera cópias locais das páginas originais, preservando figuras e fórmulas.

Uso opcional para PDFs adicionados depois: python gerar_paginas.py
"""
from pathlib import Path

import fitz
from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
STATIC = ROOT / 'static'
OUTPUT = STATIC / 'paginas'


def page_image(pdf_name, page_number):
    """Renderiza uma página referenciada por uma questão; retorna caminho em static/."""
    pdf_name = (pdf_name or '').replace('\\', '/')
    if not pdf_name.startswith('provas/') or not pdf_name.lower().endswith('.pdf'):
        return ''
    try:
        number = int(page_number)
    except (TypeError, ValueError):
        return ''
    path = (STATIC / pdf_name).resolve()
    if not path.is_relative_to((STATIC / 'provas').resolve()) or not path.is_file():
        return ''
    prefix = '__'.join(path.relative_to((STATIC / 'provas').resolve()).with_suffix('').parts)
    relative = f'paginas/{prefix}_p{number:03d}.webp'
    destination = STATIC / relative
    if destination.exists():
        try:
            with Image.open(destination) as existing:
                existing.verify()
            return relative
        except (OSError, ValueError):
            destination.unlink(missing_ok=True)
    with fitz.open(path) as document:
        if not 1 <= number <= len(document):
            return ''
        page = document[number - 1]
        bitmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        OUTPUT.mkdir(parents=True, exist_ok=True)
        image = Image.frombytes('RGB', [bitmap.width, bitmap.height], bitmap.samples)
        temporary = destination.with_suffix('.webp.tmp')
        try:
            image.save(temporary, 'WEBP', quality=84, method=5)
            with Image.open(temporary) as saved:
                saved.verify()
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
    return relative


def generate_from_workbook(file):
    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = workbook['Questoes']
        headers = [c.value for c in next(sheet.rows)]
        pairs = {(row[headers.index('arquivo_pdf')], row[headers.index('pagina')])
                 for row in sheet.iter_rows(min_row=2, values_only=True)
                 if row[headers.index('arquivo_pdf')]}
        missing = [(path, page) for path, page in sorted(pairs)
                   if not page_image(path, page)]
        if missing:
            raise ValueError(f'Páginas sem imagem: {missing}')
        return len(pairs)
    finally:
        workbook.close()


if __name__ == '__main__':
    total = generate_from_workbook(ROOT / 'imports' / 'lote1_transpetro_petrobras.xlsx')
    print(f'{total} páginas disponíveis em static/paginas')
